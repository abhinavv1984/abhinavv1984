from __future__ import annotations

import os
from typing import List, Tuple, Any
from openpyxl import load_workbook
from models import TestCase

REQUIRED_COLUMNS = [
    "TestCaseId",
    "TestCaseName",
    "Enabled",
]

OPTIONAL_COLUMNS = [
    "SqlText",
    "SqlFile",
    "Description",
    "Suite",
    "Owner",
    "PassIfZeroRows",
    "RowLimitForReport",
    "Database",
    "TimeoutSeconds",
]

ALL_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS


def _coerce_int(value: Any, default: int) -> int:
    try:
        if value is None or value == "":
            return default
        return int(value)
    except Exception:
        return default


def load_test_cases(excel_path: str) -> Tuple[List[TestCase], List[str]]:
    if not os.path.exists(excel_path):
        return [], [f"Excel not found: {excel_path}"]

    try:
        wb = load_workbook(filename=excel_path, data_only=True)
    except Exception as ex:
        return [], [f"Failed to open Excel: {ex}"]

    ws = wb.active
    headers = []
    for cell in ws[1]:
        headers.append(str(cell.value).strip() if cell.value is not None else "")

    errors: List[str] = []
    for col in REQUIRED_COLUMNS:
        if col not in headers:
            errors.append(f"Missing required column: {col}")
    if errors:
        return [], errors

    col_index = {name: idx for idx, name in enumerate(headers)}

    test_cases: List[TestCase] = []
    for row_idx in range(2, ws.max_row + 1):
        values = [ws.cell(row=row_idx, column=col + 1).value for col in range(len(headers))]
        row = {headers[i]: values[i] for i in range(len(headers))}
        try:
            tc = TestCase(
                test_case_id=str(row.get("TestCaseId", "")).strip(),
                test_case_name=str(row.get("TestCaseName", "")).strip(),
                enabled=_coerce_int(row.get("Enabled", 0), 0),
                sql_text=(str(row.get("SqlText")).strip() if row.get("SqlText") not in (None, "") else None),
                sql_file=(str(row.get("SqlFile")).strip() if row.get("SqlFile") not in (None, "") else None),
                description=(str(row.get("Description")).strip() if row.get("Description") not in (None, "") else None),
                suite=(str(row.get("Suite")).strip() if row.get("Suite") not in (None, "") else None),
                owner=(str(row.get("Owner")).strip() if row.get("Owner") not in (None, "") else None),
                pass_if_zero_rows=_coerce_int(row.get("PassIfZeroRows", 1), 1),
                row_limit_for_report=_coerce_int(row.get("RowLimitForReport", 10), 10),
                database=(str(row.get("Database")).strip() if row.get("Database") not in (None, "") else None),
                timeout_seconds=_coerce_int(row.get("TimeoutSeconds"), 0) if row.get("TimeoutSeconds") not in (None, "") else None,
            )
            test_cases.append(tc)
        except Exception as ex:
            errors.append(f"Row parse error at Excel row {row_idx}: {ex}")

    return test_cases, errors